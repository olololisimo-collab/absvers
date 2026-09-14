import docx
from docx.shared import Inches, Pt, RGBColor

doc = docx.Document()

for s in doc.sections:
    s.top_margin = Inches(0.8)
    s.bottom_margin = Inches(0.8)
    s.left_margin = Inches(0.8)
    s.right_margin = Inches(0.8)

with open('promt_magazin_abs_lokery.md', 'r', encoding='utf-8') as f:
    text = f.read()

for line in text.splitlines():
    line_str = line.strip()
    if not line_str:
        continue
    if line_str.startswith('# '):
        doc.add_heading(line_str[2:], level=1)
    elif line_str.startswith('## '):
        doc.add_heading(line_str[3:], level=2)
    elif line_str.startswith('### '):
        doc.add_heading(line_str[4:], level=3)
    elif line_str.startswith('#### '):
        doc.add_heading(line_str[5:], level=4)
    elif line_str.startswith('- ') or line_str.startswith('* '):
        doc.add_paragraph(line_str[2:], style='List Bullet')
    elif line_str.startswith('> '):
        p = doc.add_paragraph(line_str[2:])
        p.paragraph_format.left_indent = Inches(0.3)
    elif line_str.startswith('```'):
        continue
    else:
        doc.add_paragraph(line_str)

doc.save('promt_magazin_abs_lokery.docx')
print('SUCCESS')
